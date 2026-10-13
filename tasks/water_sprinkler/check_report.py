"""Проверка записки по п. 1.2 хрестоматии. Правила заданы здесь кодом, исполнитель их не задаёт и не видит.

Проверка работает по фактическим DOCX/PDF и по модели расчёта. Любое нарушение блокирует выпуск.
"""
from __future__ import annotations

import re
from pathlib import Path

import docx

from core.docs import norm, pdf_pages_text

REQUIRED_SECTIONS = [
    ("аннотация", "Аннотация"),
    ("содержание", "Содержание"),
    ("введение", "Введение"),
    ("исходные данные", "Исходные данные"),
    ("теоретическая часть", "Теоретическая часть"),
    ("расч[её]тная схема", "Расчётная (аксонометрическая) схема"),
    ("гидравлический расч[её]т", "Гидравлический расчёт"),
    ("выбор оборудования", "Выбор оборудования"),
    ("спецификаци\\w*", "Спецификация"),
    ("заключение", "Заключение"),
    ("список литературы|список использованных источников|список источников", "Список литературы"),
]
INTRO_TOPICS = [
    ("пожарн\\w* опасност", "анализ пожарной опасности объекта"),
    ("пожарн\\w* нагрузк", "пожарная нагрузка"),
    ("противопоказан|недопустим|нельзя применять|не допускается применение", "противопоказания к применению воды"),
    ("совместим", "совместимость воды с веществами и материалами"),
    ("огнетушащ\\w* способност|охлажда", "огнетушащая способность воды"),
]
FORBIDDEN = [
    "контрольная операция", "сводка независимых проверок", "машинная проверка", "внутренний журнал",
    "инструкция проверяющему", "независимая проверка", "пройдено", "статус:", "уточнить", "todo", "sha256", "хеш",
    "нейросет", "искусственн\\w* интеллект", "конвейер", "проверка результата", "монтажная ведомость",
    "перенести на план", "ведомость оросителей", "координат\\w* оросител", "lorem", "markdown", "```",
]
PAGE_MIN, PAGE_MAX = 25, 30
MAX_APPENDIX_SHARE = 0.15
MIN_CHARS_PER_PAGE = 600
MIN_CHAPTER_TAIL_CHARS = 150
MAX_TABLE_ROWS = 40
CHAPTER_START = re.compile(r"^(\d+\s+[А-ЯЁ]|введение$|заключение$|список литературы$)", re.IGNORECASE)


def _starts_chapter(page: str) -> bool:
    """Страница начинается с заголовка главы (номер страницы в колонтитуле пропускается)."""
    for line in page.splitlines():
        s = line.strip()
        if not s or s.isdigit():
            continue
        return bool(CHAPTER_START.match(s)) and len(s) < 90
    return False


def _find_page(pages: list[str], pattern: str, start: int = 0, heading_only: bool = False) -> int | None:
    for i in range(start, len(pages)):
        for line in pages[i].splitlines():
            m = re.fullmatch(rf"\s*(\d+(?:\.\d+)?\s+)?({pattern})\b(.*)", line, flags=re.IGNORECASE)
            if m and len(line) < 90 and (not heading_only or m.group(1) or not m.group(3).strip()):
                return i
    return None


def numbers_in(text: str) -> set[str]:
    return set(re.findall(r"\d+,\d+", text))


def check_report(docx_path: Path, pdf_path: Path, design: dict | None = None) -> dict:
    errors: list[str] = []
    pages = pdf_pages_text(pdf_path)
    npages = len(pages)
    low = [norm(p) for p in pages]
    full = "\n".join(low)

    if not PAGE_MIN <= npages <= PAGE_MAX:
        errors.append(f"объём {npages} с., требуется {PAGE_MIN}–{PAGE_MAX} с. (п. 1.2 хрестоматии)")

    positions = {}
    toc_page = _find_page(pages, "содержание")
    body_start = (toc_page or 0) + 1
    for pattern, title in REQUIRED_SECTIONS:
        start = 0 if title in ("Аннотация", "Содержание") else body_start
        pos = _find_page(pages, pattern, start, heading_only=True)
        if pos is None:
            errors.append(f"нет обязательного раздела «{title}»")
        positions[title] = pos
    order = [positions[t] for _, t in REQUIRED_SECTIONS if positions.get(t) is not None]
    if order != sorted(order):
        errors.append("порядок разделов не соответствует п. 1.2 хрестоматии")

    if positions.get("Введение") is not None:
        intro_end = positions.get("Исходные данные") or positions["Введение"] + 1
        intro = "\n".join(low[positions["Введение"]:intro_end + 1])
        for pattern, topic in INTRO_TOPICS:
            if not re.search(pattern, intro):
                errors.append(f"во введении не раскрыто: {topic}")

    for pat in FORBIDDEN:
        m = re.search(pat, full)
        if m:
            page = next(i for i, p in enumerate(low) if re.search(pat, p)) + 1
            errors.append(f"служебный или запрещённый текст «{m.group(0)}» (с. {page})")

    appendix = [i for i, p in enumerate(pages) if i >= body_start
                and re.search(r"^\s*приложение\s+[а-яa-z0-9]", p, flags=re.I | re.M)]
    if appendix:
        share = (npages - appendix[0]) / npages
        if share > MAX_APPENDIX_SHARE:
            errors.append(f"приложения занимают {share:.0%} объёма (с. {appendix[0] + 1}–{npages}), допустимо "
                          f"не более {MAX_APPENDIX_SHARE:.0%}")

    end = positions.get("Список литературы") or npages
    thin = []
    for i in range(2, end):
        chars = len(pages[i].strip())
        if chars >= MIN_CHARS_PER_PAGE or "рисунок" in pages[i].lower():
            continue
        chapter_tail = i + 1 < npages and _starts_chapter(pages[i + 1]) and chars >= MIN_CHAPTER_TAIL_CHARS
        if not chapter_tail:
            thin.append(i + 1)
    if thin:
        errors.append(f"почти пустые страницы основной части: {thin}")

    toc_pages = _toc_numbers(pages, toc_page)
    for title, want in toc_pages.items():
        pos = _find_page(pages, re.escape(title.lower()), body_start)
        if pos is not None and pos + 1 != want:
            errors.append(f"в содержании «{title}» указано с. {want}, фактически с. {pos + 1}")

    doc = docx.Document(str(docx_path))
    for ti, t in enumerate(doc.tables, 1):
        if len(t.rows) > MAX_TABLE_ROWS:
            errors.append(f"таблица {ti} содержит {len(t.rows)} строк – перечень по каждому элементу вместо расчёта")
    xml = doc.element.xml
    eq = xml.count("<m:oMath>") + xml.count("<m:oMath ")
    if eq < 10:
        errors.append(f"редактируемых формул {eq}, ожидается не менее 10")
    paras = [p.text.strip() for p in doc.paragraphs if len(p.text.strip()) > 120]
    dup = {p for p in paras if paras.count(p) > 1}
    if dup:
        errors.append(f"повторяющиеся абзацы: {len(dup)}")

    if design is not None:
        errors += _check_numbers(full, design)
    return dict(passed=not errors, errors=errors, pages=npages, equations=eq)


def _toc_numbers(pages: list[str], toc_page: int | None) -> dict[str, int]:
    if toc_page is None:
        return {}
    out = {}
    for line in pages[toc_page].splitlines():
        m = re.match(r"^\s*(\d+(?:\.\d+)?\s+[А-ЯЁ][^.\d]{3,80}?)[\s.]*?(\d{1,2})\s*$", line)
        if m:
            out[m.group(1).strip()] = int(m.group(2))
    return out


def _check_numbers(full: str, d: dict) -> list[str]:
    """Ключевые результаты модели обязаны присутствовать в тексте в том же виде, что и в модели."""
    def f(x, n):
        return f"{x:.{n}f}".replace(".", ",")

    expected = {
        "расчётный расход": f(d["total_flow"], 2),
        "требуемый напор": f(d["head"]["total"], 2),
        "напор у диктующего оросителя": f(d["h1"], 1),
        "потери в узле управления": f(d["head"]["huu"], 3),
    }
    errs = [f"в записке нет значения «{k}» = {v} из модели" for k, v in expected.items() if v not in full]
    if d["pump"]["name"].lower().replace(" ", "") not in full.replace(" ", ""):
        errs.append(f"в записке не указан выбранный насос {d['pump']['name']}")
    return errs
