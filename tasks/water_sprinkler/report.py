"""Сборка пояснительной записки по п. 1.2 хрестоматии.

Структура, числа, формулы, таблицы и рисунки формирует код из модели расчёта.
Sol поставляет только абзацы прозаических разделов (prose.SECTIONS).
"""
from __future__ import annotations

import math
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from core.omml import add_equation
from . import norms

FONT = "Times New Roman"

BIBLIOGRAPHY = [
    "СП 5.13130.2009. Системы противопожарной защиты. Установки пожарной сигнализации и пожаротушения "
    "автоматические. Нормы и правила проектирования. – М.: ФГУ ВНИИПО МЧС России, 2009. – 103 с.",
    "Гидравлика и противопожарное водоснабжение: учебник / под ред. Ю. А. Кошмарова. – М.: ВИПТШ МВД СССР, "
    "1985. – 285 с.",
    "Хрестоматия по контрольной работе дисциплины «Производственная и пожарная автоматика» (ЗФО). – 2016.",
    "Проектирование водяных и пенных автоматических установок пожаротушения: учеб.-метод. пособие / Л. М. Мешман, "
    "С. Г. Цариченко, В. А. Былинкин, В. В. Алешин, Р. Ю. Губин; под общ. ред. Н. П. Копылова. – М.: ВНИИПО "
    "МЧС РФ, 2002. – 413 с.",
    "Федеральный закон от 22.07.2008 № 123-ФЗ «Технический регламент о требованиях пожарной безопасности».",
    "СП 10.13130.2009. Системы противопожарной защиты. Внутренний противопожарный водопровод. Требования "
    "пожарной безопасности.",
    "ГОСТ Р 53325-2009. Техника пожарная. Технические средства пожарной автоматики. Общие технические требования. "
    "Методы испытаний.",
    norms.SPRINKLER_PASSPORT["source"] + ".",
    norms.CONTROL_VALVE_SOURCE + ".",
    "ГОСТ 10704-91. Трубы стальные электросварные прямошовные. Сортамент.",
    "ГОСТ 22247-96. Насосы центробежные консольные общего назначения для воды. Основные параметры и размеры.",
]
REF = {"sp5": 1, "koshmarov": 2, "reader": 3, "meshman": 4, "fz123": 5, "sp10": 6, "gost53325": 7,
       "passport": 8, "valve": 9, "gost10704": 10, "gost22247": 11}

HEADINGS = [
    ("Аннотация", 0), ("Введение", 0), ("1 Исходные данные", 1), ("2 Теоретическая часть", 1),
    ("3 Расстановка оросителей и расчётная схема", 1), ("4 Гидравлический расчёт", 1),
    ("5 Выбор оборудования насосной станции", 1), ("6 Спецификация оборудования", 1),
    ("Заключение", 0), ("Список литературы", 0),
]


def f(x: float, n: int = 2) -> str:
    s = f"{x:.{n}f}"
    return s.replace(".", ",")


def g(x: float) -> str:
    s = f"{x:.3f}".rstrip("0").rstrip(".")
    return s.replace(".", ",")


class Builder:
    def __init__(self, d: dict, prose: dict, meta: dict, figs: dict):
        self.d, self.prose, self.meta, self.figs = d, prose, meta, figs
        self.doc = Document()
        self.eq = 0
        self.tab = 0
        self.fig = 0
        self._setup()

    def _setup(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(14)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.first_line_indent = Cm(1.25)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size in (("Heading 1", 14), ("Heading 2", 14)):
            h = self.doc.styles[name]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = None
            rpr = h.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.append(fonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(attr), FONT)
            for t in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                fonts.attrib.pop(qn(t), None)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.first_line_indent = Cm(0) if name == "Heading 1" else Cm(1.25)
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else WD_ALIGN_PARAGRAPH.LEFT
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.5
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
        sec.different_first_page_header_footer = True
        self._page_number(sec.footer.paragraphs[0])

    def _page_number(self, p):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        run.font.size = Pt(12)

    def p(self, text: str = "", indent=True, align=None, bold=False, size=None, space_after=0, italic=False):
        para = self.doc.add_paragraph()
        if not indent:
            para.paragraph_format.first_line_indent = Cm(0)
        if align is not None:
            para.alignment = align
        if space_after:
            para.paragraph_format.space_after = Pt(space_after)
        if text:
            r = para.add_run(text)
            r.bold = bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
        return para

    def heading(self, text: str, level: int = 1, new_page: bool = False):
        h = self.doc.add_heading(text, level=level)
        if new_page:
            h.paragraph_format.page_break_before = True
        return h

    def page_break(self):
        self.p("", indent=False).paragraph_format.page_break_before = True

    def equation(self, formula: str, label: bool = True):
        para = add_equation(self.doc, formula)
        if label:
            self.eq += 1
            para.add_run(f"\t({self.eq})")
            tabs = para.paragraph_format.tab_stops
            tabs.add_tab_stop(Cm(17.0), alignment=2)
        return self.eq

    def where(self, items: list[tuple[str, str]]):
        self.p("где " + items[0][0] + " – " + items[0][1] + ";", indent=False)
        for sym, desc in items[1:]:
            self.p(f"{sym} – {desc};", indent=True)

    def prose_block(self, key: str):
        for text in self.prose[key]:
            self.p(text.strip())

    def table(self, rows: list[list[str]], widths: list[float], caption: str, header_rows: int = 1, size: int = 12):
        self.tab += 1
        cap = self.p(f"Таблица {self.tab} – {caption}", indent=False, space_after=3)
        cap.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        grid = t._tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(Cm(widths[j]).twips)))
        wide = max(range(len(widths)), key=lambda j: widths[j])
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths[j])
                para = cell.paragraphs[0]
                para.paragraph_format.first_line_indent = Cm(0)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == wide and i >= header_rows) else WD_ALIGN_PARAGRAPH.CENTER
                r = para.add_run(val)
                r.font.size = Pt(size)
                r.bold = i < header_rows
        self.p("")
        return t

    def figure(self, path: Path, caption: str, width_cm: float = 16.0):
        self.fig += 1
        para = self.p(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        self.p(f"Рисунок {self.fig} – {caption}", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
        return self.fig

    # ---------------------------------------------------------------- разделы
    def title_page(self):
        m, v = self.meta, self.d["variant"]
        c = WD_ALIGN_PARAGRAPH.CENTER
        for line in m.get("institution", []):
            self.p(line, indent=False, align=c)
        for _ in range(5):
            self.p("", indent=False)
        self.p("КОНТРОЛЬНАЯ РАБОТА", indent=False, align=c, bold=True, size=16)
        self.p("по дисциплине «Производственная и пожарная автоматика»", indent=False, align=c)
        self.p(f"на тему: «Расчёт спринклерной установки водяного пожаротушения. {v['room']}»", indent=False, align=c)
        self.p(f"Вариант {v['number']}", indent=False, align=c)
        for _ in range(5):
            self.p("", indent=False)
        r = WD_ALIGN_PARAGRAPH.RIGHT
        self.p(f"Выполнил: студент группы {m.get('group') or '__________'}", indent=False, align=r)
        self.p(m.get("student") or "______________________", indent=False, align=r)
        self.p(f"Проверил: {m.get('teacher') or '______________________'}", indent=False, align=r)
        for _ in range(4):
            self.p("", indent=False)
        self.p(f"{m.get('city', '')} {m.get('year', '')}".strip(), indent=False, align=c)

    def task_page(self):
        d, v, p = self.d, self.d["variant"], self.d["params"]
        self.page_break()
        self.p("ЗАДАНИЕ НА КОНТРОЛЬНУЮ РАБОТУ", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=6)
        self.p(f"Рассчитать параметры спринклерной установки водяного пожаротушения для помещения «{v['room']}» "
               f"(вариант {v['number']} по таблице 1.1 [{REF['reader']}]), выполнить гидравлический расчёт, "
               "определить требуемый напор насосной станции, подобрать оборудование и выполнить графическую часть.")
        rows = [["Параметр", "Значение"],
                ["Номер варианта", str(v["number"])],
                ["Тип помещения", v["room"]],
                ["Размеры помещения (длина × ширина × высота), м", f"{g(v['length'])} × {g(v['width'])} × {g(v['height'])}"],
                ["Тип оросителя", d["sprinkler"]["name"]]]
        self.table(rows, [10.5, 6.0], "Исходные данные варианта")
        self.p("Состав работы: расчётно-пояснительная записка объёмом 25–30 страниц формата А4; графическая часть – "
               "три листа формата А3: план помещения с трубной разводкой и расстановкой оросителей, аксонометрическая "
               "схема установки, расчётная схема с таблицей результатов гидравлического расчёта.")

    def annotation(self):
        self.heading("Аннотация", new_page=True)
        self.prose_block("annotation")
        self.p("Пояснительная записка содержит {PAGES} с., {FIGS} рис., {TABS} табл., "
               f"{len(BIBLIOGRAPHY)} источников. Графическая часть выполнена на трёх листах формата А3.")

    def toc(self, pages: dict[str, int]):
        self.heading("Содержание", new_page=True)
        for title, _level in HEADINGS:
            if title == "Аннотация":
                continue
            para = self.p(indent=False)
            para.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), alignment=2, leader=1)
            para.add_run(f"{title}\t{pages.get(title, '')}")
            for sub in self.subheadings.get(title, []):
                sp = self.p(indent=False)
                sp.paragraph_format.left_indent = Cm(0.6)
                sp.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), alignment=2, leader=1)
                sp.add_run(f"{sub}\t{pages.get(sub, '')}")

    def introduction(self):
        self.heading("Введение", new_page=True)
        self.prose_block("introduction")

    def source_data(self):
        d, v, p, spr = self.d, self.d["variant"], self.d["params"], self.d["sprinkler"]
        self.heading("1 Исходные данные", new_page=True)
        self.sub("1.1 Характеристика защищаемого объекта")
        self.prose_block("object")
        self.sub("1.2 Параметры установки пожаротушения")
        self.p(f"По приложению Б [{REF['sp5']}] помещение относится к группе {p['group']} по степени опасности "
               f"развития пожара. Параметры установки для этой группы приняты по таблице 5.1 [{REF['sp5']}] "
               f"и сведены в таблицу {self.tab + 1}.")
        rows = [["Параметр", "Обозначение", "Значение"],
                ["Группа помещения по приложению Б [1]", "–", p["group"]],
                ["Интенсивность орошения, л/(с·м²)", "I", g(p["intensity"])],
                ["Минимальная площадь для расчёта расхода, м²", "S", g(p["area"])],
                ["Нормативный расход, л/с, не менее", "Qн", g(p["flow"])],
                ["Продолжительность работы, мин", "t", str(p["duration"])],
                ["Максимальное расстояние между оросителями, м", "L", g(p["spacing"])],
                ["Максимальное расстояние от оросителя до стены, м", "L/2", g(p["spacing"] / 2)],
                ["Тип оросителя", "–", spr["name"]],
                ["Коэффициент производительности оросителя [8]", "k", g(spr["K"])],
                ["Защищаемая оросителем площадь по паспорту [8], м²", "Fор", g(spr["passport"]["protected_area"])],
                ["Диапазон рабочего давления по паспорту [8], МПа", "–",
                 f"{g(spr['passport']['p_min_mpa'])}–{g(spr['passport']['p_max_mpa'])}"]]
        self.table(rows, [10.0, 2.5, 4.0], "Исходные данные для расчёта")
        self.p(f"Коэффициент производительности k = {g(spr['K'])} в паспорте [{REF['passport']}] приведён в "
               "дм³/(с·МПа⁰·⁵). При записи напора в метрах водяного столба, как в методике [3], формула расхода "
               "сохраняет вид q = k·√H, так как 10·k·√(H/100) = k·√H.")
        self.sub("1.3 Основные расчётные формулы")
        self.p(f"Гидравлический расчёт выполняется по методике приложения В [{REF['sp5']}] в форме, принятой в [{REF['reader']}]. "
               "Расход оросителя определяется по формуле")
        self.equation("q=k\\sqrt{H}")
        self.where([("q", "расход оросителя, л/с"), ("k", "коэффициент производительности оросителя"),
                    ("H", "напор перед оросителем, м вод. ст.")])
        self.p("Потери напора на участке трубопровода определяются по формуле")
        self.equation("ΔH=\\frac{\\sup{Q}{2}·l}{\\sub{K}{т}}")
        self.where([("Q", "расход воды на участке, л/с"), ("l", "длина участка, м"),
                    ("Kт", f"удельная характеристика трубопровода по таблице В.2 [{REF['sp5']}], л⁶/с²")])
        self.p("Расход каждого последующего рядка, выполненного конструктивно одинаково с первым, определяется через "
               "обобщённую характеристику рядка")
        self.equation("B=\\frac{\\sup{\\sub{Q}{I}}{2}}{\\sub{H}{a}},  \\sub{Q}{II}=\\sqrt{B·\\sub{H}{b}}")
        self.where([("QI", "расход первого рядка, л/с"), ("Ha", "напор в точке подключения первого рядка, м"),
                    ("Hb", "напор в точке подключения второго рядка, м")])
        self.p("Требуемый напор насосной станции определяется по формуле")
        self.eq_head = self.equation("\\sub{H}{тр}=\\sub{H}{1}+ΣΔH+\\sub{H}{м}+\\sub{H}{уу}+Z")
        self.where([("H1", "напор у диктующего оросителя, м"), ("ΣΔH", "сумма потерь напора по длине, м"),
                    ("Hм", f"потери в местных сопротивлениях, принимаемые 20 % от ΣΔH по В.3.14 [{REF['sp5']}], м"),
                    ("Hуу", "потери напора в узле управления, м"),
                    ("Z", "геометрическая высота диктующего оросителя над осью насоса, м")])
        self.p(f"Потери напора в узле управления определяются по В.3.13 [{REF['sp5']}]")
        self.equation("\\sub{H}{уу}=\\sub{ξ}{уу}·γ·\\sup{Q}{2}")
        self.where([("ξуу", f"коэффициент потерь давления узла управления по технической документации [{REF['valve']}]"),
                    ("γ", "плотность воды, 1000 кг/м³"), ("Q", "расход через узел управления, м³/ч")])

    def theory(self):
        self.heading("2 Теоретическая часть", new_page=True)
        self.sub("2.1 Назначение автоматической установки пожаротушения")
        self.prose_block("purpose")
        self.sub("2.2 Состав установки")
        self.prose_block("composition")
        self.sub("2.3 Принцип работы установки")
        self.prose_block("operation")

    def layout(self):
        d, g_ = self.d, self.d["grid"]
        p = d["params"]
        self.heading("3 Расстановка оросителей и расчётная схема", new_page=True)
        self.sub("3.1 Расстановка оросителей")
        self.prose_block("layout")
        self.p(f"Минимальное число оросителей, обеспечивающих расчётную площадь, определяется по В.3.2 [{REF['sp5']}]:")
        self.equation(f"n≥\\frac{{S}}{{Ω}}=\\frac{{{g(p['area'])}}}{{{f(g_['a'])}·{f(g_['b'])}}}="
                      f"{f(p['area'] / g_['omega'])}",)
        self.p(f"Принято {d['zone']['n_need']} оросителей; диктующая зона включает {d['zone']['rows']} рядка по "
               f"{g_['n_r']} оросителей, всего {d['zone']['count']} оросителей на площади {f(d['zone']['area'], 1)} м², "
               f"что не меньше нормативной площади {g(p['area'])} м². Площадь, защищаемая одним оросителем, составляет "
               f"{f(g_['omega'])} м², что не превышает паспортного значения {g(d['sprinkler']['passport']['protected_area'])} м². "
               f"План расстановки оросителей показан на рисунке {self.fig + 1}.")
        self.figure(self.figs["plan"], "План помещения с расстановкой оросителей (зачернены оросители диктующей зоны)")
        self.sub("3.2 Расчётная схема")
        self.p(f"Для гидравлического расчёта выполнена аксонометрическая расчётная схема диктующей зоны (рисунок {self.fig + 1}). "
               "Точка 1 – диктующий ороситель, наиболее удалённый от узла управления. Номера точек схемы соответствуют "
               "участкам таблицы результатов гидравлического расчёта в разделе 4.5. Диаметры участков приняты "
               "по результатам подбора в разделе 4.1.")
        self.figure(self.figs["scheme"], "Расчётная схема диктующей зоны", 15.5)

    def hydraulics(self):
        d = self.d
        sec = d["sections"]
        k, h1 = d["sprinkler"]["K"], d["h1"]
        g_ = d["grid"]
        per_side = g_["per_side"]
        self.heading("4 Гидравлический расчёт", new_page=True)
        self.sub("4.1 Подбор диаметров трубопроводов")
        self.p(f"Ориентировочный диаметр распределительного трубопровода по таблице В.3 [{REF['sp5']}] для "
               f"{per_side} оросителей на ветви составляет Ду{d['dn_ref_b3']}. Для подбора оптимального сочетания "
               "выполнен расчёт для нескольких диаметров распределительного и питающего трубопроводов; результаты "
               f"приведены в таблице {self.tab + 1}. Оптимальным принято сочетание, при котором скорость воды не превышает "
               f"{g(norms.VELOCITY_LIMIT)} м/с, а дальнейшее увеличение диаметра питающего трубопровода снижает "
               f"требуемый напор менее чем на {int(norms.HEAD_GAIN_THRESHOLD * 100)} %.")
        rows = [["Распредели-\nтельный", "Питающий", "H₁, м", "Q, л/с", "Нтр, м", "vmax, м/с", "Решение"]]
        for c in d["candidates"]:
            rows.append([f"Ду{c['dn_dist']}", f"Ду{c['dn_feed']}", f(c["h1"], 1), f(c["flow"]), f(c["h_req"]),
                         f(max(c["v_dist"], c["v_feed"])), "принято" if c["chosen"] else "–"])
        self.table(rows, [2.6, 2.4, 1.8, 2.0, 2.2, 2.3, 2.8], "Сравнение вариантов диаметров трубопроводов")
        self.p(f"Принят распределительный трубопровод Ду{d['chosen']['dn_dist']} ({sec[0]['pipe']} мм) и питающий "
               f"трубопровод Ду{d['chosen']['dn_feed']} ({sec[-1]['pipe']} мм) по ГОСТ 10704-91 [{REF['gost10704']}].")

        self.sub("4.2 Определение напора у диктующего оросителя")
        q_need = d["q1_required"]
        self.p("Расход диктующего оросителя должен обеспечивать нормативную интенсивность орошения на площади, "
               "защищаемой одним оросителем:")
        self.equation(f"\\sub{{q}}{{min}}=I·Ω={g(d['params']['intensity'])}·{f(g_['omega'])}={f(q_need, 3)} л/с")
        self.equation(f"\\sub{{H}}{{min}}=\\sup{{(\\frac{{\\sub{{q}}{{min}}}}{{k}})}}{{2}}=\\sup{{(\\frac{{{f(q_need, 3)}}}{{{g(k)}}})}}{{2}}="
                      f"{f((q_need / k) ** 2)} м")
        self.p(f"Минимальное рабочее давление оросителя по паспорту [{REF['passport']}] – "
               f"{g(d['sprinkler']['passport']['p_min_mpa'])} МПа, что соответствует 5 м вод. ст. Кроме того, суммарный "
               f"расход диктующей зоны должен быть не менее нормативного {g(d['params']['flow'])} л/с. Напор у диктующего "
               f"оросителя подобран последовательными расчётами так, чтобы выполнялись все три условия, и принят "
               f"H₁ = {f(h1, 1)} м вод. ст. = {f(h1 / 100, 3)} МПа.")

        self.sub("4.3 Расчёт распределительного трубопровода")
        heads, flows = d["branch"]["heads"], d["branch"]["flows"]
        cum = 0.0
        for i in range(per_side):
            s = sec[i]
            self.p(f"Расход {i + 1}-го оросителя:", indent=True)
            self.equation(f"\\sub{{q}}{{{i + 1}}}={g(k)}·\\sqrt{{{f(heads[i], 3)}}}={f(flows[i], 3)} л/с")
            cum += flows[i]
            self.p(f"Потери напора на участке {s['name']} (расход {f(s['flow'], 3)} л/с, длина {f(s['length'])} м):")
            self.equation(f"\\sub{{ΔH}}{{{s['name']}}}=\\frac{{\\sup{{{f(s['flow'], 3)}}}{{2}}·{f(s['length'])}}}"
                          f"{{{g(norms.PIPES[s['dn']][2])}}}={f(s['loss'], 3)} м")
            nxt = heads[i] + s["loss"]
            self.p(f"Напор в точке {i + 2}:")
            self.equation(f"\\sub{{H}}{{{i + 2}}}={f(heads[i], 3)}+{f(s['loss'], 3)}={f(nxt, 3)} м")
        q_branch = sum(flows)
        q_row = d["row_flows"][0]
        self.p(f"Расход одной ветви рядка I составляет {f(q_branch, 3)} л/с. Ветви рядка выполнены симметрично, "
               f"поэтому расход рядка I равен QI = 2·{f(q_branch, 3)} = {f(q_row, 3)} л/с при напоре в точке "
               f"{per_side + 1} Ha = {f(d['row_heads'][0], 3)} м.")
        self.sub("4.4 Расчёт питающего трубопровода")
        self.p("Обобщённая характеристика рядка I:")
        self.equation(f"B=\\frac{{\\sup{{{f(q_row, 3)}}}{{2}}}}{{{f(d['row_heads'][0], 3)}}}={f(d['b_char'], 3)}")
        total = q_row
        feed_secs = sec[per_side:]
        for r in range(1, d["zone"]["rows"]):
            s = feed_secs[r - 1]
            self.p(f"Потери напора на участке {s['name']} питающего трубопровода при расходе {f(s['flow'], 3)} л/с:")
            self.equation(f"\\sub{{ΔH}}{{{s['name']}}}=\\frac{{\\sup{{{f(s['flow'], 3)}}}{{2}}·{f(s['length'])}}}"
                          f"{{{g(norms.PIPES[s['dn']][2])}}}={f(s['loss'], 3)} м")
            self.p(f"Напор в точке подключения рядка {self.roman(r + 1)} и его расход:")
            self.equation(f"\\sub{{H}}{{{per_side + 1 + r}}}={f(d['row_heads'][r - 1], 3)}+{f(s['loss'], 3)}="
                          f"{f(d['row_heads'][r], 3)} м")
            self.equation(f"\\sub{{Q}}{{{self.roman(r + 1)}}}=\\sqrt{{{f(d['b_char'], 3)}·{f(d['row_heads'][r], 3)}}}="
                          f"{f(d['row_flows'][r], 3)} л/с")
            total += d["row_flows"][r]
        self.p(f"Суммарный расход диктующей зоны: Q = {' + '.join(f(x, 3) for x in d['row_flows'])} = "
               f"{f(d['total_flow'], 3)} л/с, что не меньше нормативного расхода {g(d['params']['flow'])} л/с.")
        for s in feed_secs[d["zone"]["rows"] - 1:]:
            self.p(f"Потери напора на участке {s['name']} ({s['kind'].lower()} трубопровод, длина {f(s['length'])} м):")
            self.equation(f"\\sub{{ΔH}}{{{s['name']}}}=\\frac{{\\sup{{{f(s['flow'], 3)}}}{{2}}·{f(s['length'])}}}"
                          f"{{{g(norms.PIPES[s['dn']][2])}}}={f(s['loss'], 3)} м")

        self.sub("4.5 Определение требуемого напора")
        hd, vv = d["head"], d["valve"]
        self.p(f"Потери напора в узле управления {vv['name']} при расходе {f(vv['flow_m3h'], 1)} м³/ч "
               f"(ξуу = {f(vv['xi'] * 1e7, 2)}·10⁻⁷ по [{REF['valve']}]):")
        self.equation(f"\\sub{{H}}{{уу}}={f(vv['xi'] * 1e7, 2)}·\\sup{{10}}{{-7}}·1000·\\sup{{{f(vv['flow_m3h'], 1)}}}{{2}}="
                      f"{f(hd['huu'], 3)} м")
        self.p(f"Требуемый напор насосной станции по формуле ({self.eq_head}):")
        self.equation(f"\\sub{{H}}{{тр}}={f(hd['h1'], 1)}+{f(hd['sum_loss'], 3)}+0,2·{f(hd['sum_loss'], 3)}+"
                      f"{f(hd['huu'], 3)}+{g(hd['z'])}={f(hd['total'], 2)} м")
        self.p(f"Результаты гидравлического расчёта сведены в таблицу {self.tab + 1}, оформленную по примеру [3].")
        self.result_table()
        self.p(f"Общий расход воды на спринклерное пожаротушение: Qобщ = {f(d['total_flow'], 2)} л/с = "
               f"{f(d['total_flow_m3h'], 1)} м³/ч. Требуемый напор: Нтр = {f(hd['total'], 2)} м вод. ст. = "
               f"{f(hd['total'] / 100, 3)} МПа. Объём воды на {d['params']['duration']} мин работы установки составляет "
               f"{f(d['water_volume'], 1)} м³.")

    def result_table(self):
        d, hd = self.d, self.d["head"]
        rows = [["№ п/п", "Наименование трубопровода", "Расход Q, л/с", "Труба по ГОСТ 10704-91", "Длина L, м", "Потери H, м"]]
        for i, s in enumerate(d["sections"], 1):
            rows.append([str(i), f"{s['kind']} трубопровод. Участок {s['name']}", f(s["flow"], 3), s["pipe"],
                         f(s["length"]), f(s["loss"], 3)])
        n = len(d["sections"])
        rows += [["", "Итого потерь по расчёту", "", "", "", f(hd["sum_loss"], 3)],
                 [str(n + 1), "С учётом потерь в местных сопротивлениях (+20 %)", "", "", "", f(hd["sum_loss"] * 1.2, 3)],
                 [str(n + 2), f"Потери напора в узле управления. Точка {d['sections'][-1]['name'].split('-')[1]}", "", "", "",
                  f(hd["huu"], 3)],
                 [str(n + 3), "Разность гидравлических отметок", "", "", "", g(hd["z"])],
                 [str(n + 4), "Напор у диктующего оросителя", "", "", "", f(hd["h1"], 1)],
                 ["", "Итого требуемый напор", "", "", "", f(hd["total"], 2)]]
        self.table(rows, [1.2, 6.6, 2.2, 2.6, 1.8, 2.2], "Результаты гидравлического расчёта", size=11)

    def equipment(self):
        d, pump = self.d, self.d["pump"]
        self.heading("5 Выбор оборудования насосной станции", new_page=True)
        self.prose_block("equipment")
        self.p(f"Насос подбирается по требуемой подаче Q = {f(pump['q_req_m3h'], 1)} м³/ч и требуемому напору "
               f"Н = {f(pump['h_req'], 2)} м. Принят консольный насос {pump['name']} [{REF['gost22247']}] с номинальной "
               f"подачей {g(pump['q_nom'])} м³/ч и напором {g(pump['h_nom'])} м, мощность электродвигателя "
               f"{g(pump['power'])} кВт. Запас по подаче составляет {f((pump['q_nom'] / pump['q_req_m3h'] - 1) * 100, 0)} %, "
               f"по напору – {f((pump['h_nom'] / pump['h_req'] - 1) * 100, 0)} %. Положение рабочей точки относительно "
               f"характеристики сети показано на рисунке {self.fig + 1}.")
        self.figure(self.figs["pump"], "Характеристика сети и номинальный режим выбранного насоса", 15.0)
        self.p("Насосная станция включает рабочий и резервный пожарные насосы одной марки и жокей-насос для "
               f"поддержания давления в дежурном режиме. Узел управления – {d['valve']['name']} [{REF['valve']}].")

    def specification(self):
        self.heading("6 Спецификация оборудования", new_page=True)
        rows = [["Поз.", "Наименование", "Обозначение", "Ед.", "Кол."]]
        for r in self.d["specification"]:
            qty = g(r["qty"]) if isinstance(r["qty"], float) else str(r["qty"])
            rows.append([str(r["pos"]), r["name"], r["mark"], r["unit"], qty])
        self.table(rows, [1.2, 8.4, 4.0, 1.3, 1.6], "Спецификация оборудования и материалов", size=11)

    def conclusion(self):
        self.heading("Заключение", new_page=True)
        self.prose_block("conclusion")

    def bibliography(self):
        self.heading("Список литературы", new_page=True)
        for i, b in enumerate(BIBLIOGRAPHY, 1):
            self.p(f"{i}. {b}")

    def sub(self, title: str):
        self.subheadings.setdefault(self.current, []).append(title)
        return self.doc.add_heading(title, level=2)

    @staticmethod
    def roman(n: int) -> str:
        return ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"][n]


def build_report(d: dict, prose: dict, meta: dict, figs: dict, out: Path, pages: dict | None = None,
                 counts: dict | None = None) -> dict:
    """Строит DOCX. Содержание вставляется после аннотации; номера страниц подставляются во втором проходе."""
    b = Builder(d, prose, meta, figs)
    b.subheadings, b.current = {}, None
    b.title_page()
    b.task_page()
    b.current = "Аннотация"
    b.annotation()
    b.subheadings_cache = None
    toc_pos = len(b.doc.element.body) - (1 if b.doc.element.body[-1].tag == qn("w:sectPr") else 0)
    steps = [("Введение", b.introduction), ("1 Исходные данные", b.source_data),
             ("2 Теоретическая часть", b.theory), ("3 Расстановка оросителей и расчётная схема", b.layout),
             ("4 Гидравлический расчёт", b.hydraulics), ("5 Выбор оборудования насосной станции", b.equipment),
             ("6 Спецификация оборудования", b.specification), ("Заключение", b.conclusion),
             ("Список литературы", b.bibliography)]
    for title, fn in steps:
        b.current = title
        fn()
    body = b.doc.element.body
    tail = list(body)[toc_pos:]
    for el in tail:
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    b.toc(pages or {})
    for el in tail:
        if el.tag != qn("w:sectPr"):
            body.insert(len(body) - 1, el)
    counts = counts or {}
    for para in b.doc.paragraphs:
        if "{PAGES}" in para.text:
            text = para.text.replace("{PAGES}", str(counts.get("pages", ""))).replace("{FIGS}", str(b.fig)) \
                .replace("{TABS}", str(b.tab))
            for r in para.runs[1:]:
                r.text = ""
            para.runs[0].text = text
    b.doc.save(str(out))
    all_headings = []
    for title, _ in HEADINGS:
        all_headings.append(title)
        all_headings += b.subheadings.get(title, [])
    return dict(headings=all_headings, equations=b.eq, tables=b.tab, figures=b.fig)
